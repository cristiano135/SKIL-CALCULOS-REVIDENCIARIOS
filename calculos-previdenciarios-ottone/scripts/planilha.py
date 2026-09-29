"""Gera a planilha .xlsx no modelo do escritório (Resumo, Evolução da RMI, Diferenças, Honorários...).

As contas ficam em fórmulas: alterar RMI, valores recebidos, índices ou % de honorários na própria
planilha recalcula tudo. Mudanças de datas (DIB, prescrição, período) exigem rodar o script de novo.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from motor import amd, c_fmt, d_fmt

AZUL = "1F3864"
CINZA = "D9E1F2"
MOEDA = '#,##0.00'
TITULO = Font(bold=True, size=14, color=AZUL)
CAB = Font(bold=True, color="FFFFFF")
CAB_FILL = PatternFill("solid", fgColor=AZUL)
TOT_FILL = PatternFill("solid", fgColor=CINZA)
ENTRADA = PatternFill("solid", fgColor="FFF2CC")   # células editáveis
FINO = Side(style="thin", color="A6A6A6")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _cab(ws, row, titulos, larguras=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = CAB, CAB_FILL, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for i, w in enumerate(larguras or [], 1):
        ws.column_dimensions[L(i)].width = w


def _cel(ws, row, col, valor, fmt=None, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=valor)
    c.border = BORDA
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = Font(bold=True)
    if fill:
        c.fill = fill
    return c


def _identificacao(ws, caso, titulo, p):
    seg, ben, proc = caso["segurado"], caso.get("beneficio", {}), caso.get("processo", {})
    ws["A1"] = titulo
    ws["A1"].font = TITULO
    linhas = [
        ("Nome do Segurado:", seg.get("nome", ""), "NIT:", seg.get("nit", "")),
        ("CPF:", seg.get("cpf", ""), "Data de Nascimento:", seg.get("nascimento", "")),
        ("Sexo:", {"M": "Masculino", "F": "Feminino"}.get(seg.get("sexo", "")[:1].upper(), ""),
         "Nº do Benefício:", ben.get("nb", "")),
        ("Espécie do Benefício:", ben.get("especie", ""), "Nº do Processo:", proc.get("numero", "")),
    ]
    if p:
        linhas.append(("Início do Benefício (DIB):", d_fmt(p["dib"]), "Data da Atualização:", d_fmt(p["atual"])))
        linhas.append(("Ajuizamento:", d_fmt(p["ajuizamento"]), "Data da Prescrição:", d_fmt(p["prescricao"])))
    for i, (a, b, c, d_) in enumerate(linhas, 2):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True)
        ws.cell(row=i, column=2, value=b)
        ws.cell(row=i, column=4, value=c).font = Font(bold=True)
        ws.cell(row=i, column=5, value=d_)
    return len(linhas) + 3


def _impressao(ws, freeze=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if freeze:
        ws.freeze_panes = freeze


def gerar(caminho, caso, ben=None, liq=None, series=None, avisos=()):
    wb = Workbook()
    wb.remove(wb.active)
    resumo = wb.create_sheet("Resumo")
    par = wb.create_sheet("Parâmetros")
    P = _parametros(par, caso, ben, liq)
    if ben:
        _tempo(wb.create_sheet("Tempo de Contribuição"), caso, ben)
        _regras(wb.create_sheet("Regras EC 103"), caso, ben)
        _pbc(wb.create_sheet("Salários e RMI"), caso, ben)
    if ben is not None or caso.get("vinculos"):
        _sem_salario(wb.create_sheet("Competências s-SC"), caso, ben)
    if liq:
        ev_rows = _evolucao(wb.create_sheet("Evolução RMI"), caso, liq, P)
        dif = _diferencas(wb.create_sheet("Diferenças"), caso, liq, P, ev_rows)
        hon = _honorarios(wb.create_sheet("Honorários"), caso, liq, P, dif)
        _ir(wb.create_sheet("IR (RRA)"), caso, liq, dif)
        _resumo(resumo, caso, liq, P, dif, hon)
    else:
        _resumo_beneficio(resumo, caso, ben)
    if series is not None:
        _indices(wb.create_sheet("Índices"), series)
    if avisos:
        ws = wb.create_sheet("Avisos")
        ws["A1"] = "Pontos a conferir antes de usar o cálculo"
        ws["A1"].font = TITULO
        for i, a in enumerate(avisos, 3):
            ws.cell(row=i, column=1, value="• " + a)
        ws.column_dimensions["A"].width = 140
    wb.save(caminho)


# --------------------------------------------------------------------------- parâmetros

def _parametros(ws, caso, ben, liq):
    ws["A1"] = "Parâmetros do cálculo (células amarelas podem ser editadas)"
    ws["A1"].font = TITULO
    seg, b, proc, lq = caso["segurado"], caso.get("beneficio", {}), caso.get("processo", {}), caso.get("liquidacao", {})
    adv = caso.get("advogado", {})
    itens = [
        ("nome", "Nome do segurado", seg.get("nome", ""), None, False),
        ("nit", "NIT", seg.get("nit", ""), None, False),
        ("cpf", "CPF", seg.get("cpf", ""), None, False),
        ("nasc", "Data de nascimento", seg.get("nascimento", ""), None, False),
        ("sexo", "Sexo", seg.get("sexo", ""), None, False),
        ("nb", "Nº do benefício", b.get("nb", ""), None, False),
        ("especie", "Espécie", b.get("especie", ""), None, False),
        ("der", "DER", b.get("der", ""), None, False),
    ]
    if liq:
        itens += [
            ("dib", "DIB", d_fmt(liq["dib"]), None, False),
            ("rmi", "RMI (renda mensal inicial)", liq["rmi"], MOEDA, True),
            ("processo", "Nº do processo", proc.get("numero", ""), None, False),
            ("inscricao", "Inscrição no INSS", proc.get("inscricao_inss", ""), None, False),
            ("ajuiz", "Ajuizamento", d_fmt(liq["ajuizamento"]), None, False),
            ("citacao", "Citação (início dos juros de mora)", d_fmt(liq["citacao"]), None, False),
            ("prescr", "Data da prescrição quinquenal", d_fmt(liq["prescricao"]), None, False),
            ("atual", "Data da atualização", d_fmt(liq["atual"]), None, False),
            ("de", "Diferenças desde", d_fmt(liq["inicio"]), None, False),
            ("ate", "Diferenças até", d_fmt(liq["fim"]), None, False),
            ("pct", "Honorários (%)", liq["pct"], "0.0000", True),
            ("hon_ate", "Honorários: parcelas até (Súmula 111/STJ)", lq.get("honorarios_ate") or "todas", None, False),
            ("hon_fixo", "Honorários: valor fixo", liq["honorarios"]["fixo"], MOEDA, True),
            ("acordo", "(-) Acordo/Deságio", liq["acordo"], MOEDA, True),
            ("multa", "Multa", liq["multa"], MOEDA, True),
            ("dano", "Dano moral/material", liq["dano"], MOEDA, True),
            ("sm", "Salário mínimo na data da atualização", liq["causa"]["salario_minimo"], MOEDA, True),
            ("criterio", "Índice de correção", liq["criterio"], None, False),
        ]
    itens += [("adv", "Responsável pelo cálculo", adv.get("nome", ""), None, False),
              ("oab", "OAB", adv.get("oab", ""), None, False)]
    P = {}
    for i, (k, rot, v, fmt, edit) in enumerate(itens, 3):
        ws.cell(row=i, column=1, value=rot).font = Font(bold=True)
        c = _cel(ws, i, 2, v, fmt, fill=ENTRADA if edit else None)
        P[k] = f"'Parâmetros'!$B${i}"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 48
    return P


# --------------------------------------------------------------------------- benefício

def _tempo(ws, caso, ben):
    r0 = _identificacao(ws, caso, "Tempo de Contribuição e Carência", None)
    _cab(ws, r0, ["Seq.", "Empregador / Origem", "Tipo", "Início", "Fim", "Especial", "Tempo (a/m/d)",
                  "Carência?", "Computado?"], [6, 42, 26, 12, 12, 10, 16, 10, 11])
    from motor import d, dias360
    r = r0 + 1
    for v in caso.get("vinculos", []):
        dias = dias360(d(v["inicio"]), d(v["fim"]))
        vals = [v.get("seq"), v.get("empregador", ""), v.get("tipo", ""), v["inicio"], v["fim"],
                f"{v['especial']} anos" if v.get("especial") else "", amd(dias),
                "sim" if v.get("carencia", True) else "não", "sim" if v.get("computar", True) else "não"]
        for j, x in enumerate(vals, 1):
            _cel(ws, r, j, x)
        r += 1
    r += 1
    for rot, val in [("Tempo até 13/11/2019 (com conversão do especial)", ben["tempo_ate_reforma"]),
                     ("Tempo até a DER (com conversão até 13/11/2019)", ben["tempo_ate_der"]),
                     ("Tempo até a DER (sem conversão)", ben["tempo_sem_conversao_der"]),
                     ("Tempo especial até a DER", ben["tempo_especial_der"]),
                     ("Carência até a DER (meses)", ben["carencia_der"]),
                     ("Idade na DER", ben["idade_der"])]:
        ws.cell(row=r, column=2, value=rot).font = Font(bold=True)
        _cel(ws, r, 7, val, fill=TOT_FILL)
        r += 1
    ws.cell(row=r + 1, column=2, value="Concomitâncias contadas uma vez. Conversão especial→comum só até "
                                       "13/11/2019 (EC 103/2019, art. 25, §2º). Contagem em ano de 360 / mês de 30 dias.")
    _impressao(ws, f"A{r0 + 1}")


def _regras(ws, caso, ben):
    r0 = _identificacao(ws, caso, "Regras de Aposentadoria (direito adquirido, transições da EC 103/2019 e regra permanente)", None)
    _cab(ws, r0, ["Regra", "Fundamento", "Cumpre na DER?", "Data em que cumpre", "Requisitos na DER",
                  "Requisitos na data", "DIB considerada", "Média", "Coef. (%)", "Fator prev.", "RMI"],
         [40, 26, 11, 13, 55, 55, 12, 12, 9, 10, 12])
    r = r0 + 1
    for x in ben["regras"]:
        vals = [x["regra"], x["fundamento"], "SIM" if x["cumpre_na_der"] else "não", x["data_cumprimento"] or "—",
                x["requisitos_na_der"], x["requisitos_na_data"], x.get("dib", ""), x.get("media"),
                x.get("coeficiente"), x.get("fator_previdenciario"), x.get("rmi")]
        for j, v in enumerate(vals, 1):
            c = _cel(ws, r, j, v, MOEDA if j in (8, 11) else None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if x["regra"] == ben["melhor_regra"]:
                c.fill = TOT_FILL
        r += 1
    ws.cell(row=r + 1, column=1, value=f"Regra mais vantajosa: {ben['melhor_regra']}").font = Font(bold=True)
    ws.cell(row=r + 2, column=1, value="Datas posteriores à DER supõem contribuições contínuas após o último vínculo "
                                       "(reafirmação da DER, Tema 995/STJ).")
    _impressao(ws, f"B{r0 + 1}")


def _pbc(ws, caso, ben):
    r0 = _identificacao(ws, caso, "Salários de Contribuição (PBC desde 07/1994) e Média", None)
    _cab(ws, r0, ["Competência", "Salário (CNIS)", "Teto", "Salário limitado", "Fator de correção",
                  "Salário corrigido"], [13, 15, 12, 15, 15, 16])
    r = r0 + 1
    for x in ben["pbc"]:
        _cel(ws, r, 1, x["competencia"])
        _cel(ws, r, 2, x["valor"], MOEDA)
        _cel(ws, r, 3, x["teto"], MOEDA)
        _cel(ws, r, 4, f"=IF(C{r}=\"\",B{r},MIN(B{r},C{r}))", MOEDA)
        _cel(ws, r, 5, x["fator"], "0.000000", fill=ENTRADA)
        _cel(ws, r, 6, f"=ROUND(D{r}*E{r},2)" if x["fator"] is not None else x["corrigido"], MOEDA)
        r += 1
    ult = r - 1
    ws.cell(row=r + 1, column=1, value="Média de 100% dos salários (art. 26 da EC 103/2019):").font = Font(bold=True)
    _cel(ws, r + 1, 6, f"=IFERROR(ROUND(AVERAGE(F{r0 + 1}:F{ult}),2),0)", MOEDA, bold=True, fill=TOT_FILL)
    ws.cell(row=r + 2, column=1, value="Quantidade de salários:").font = Font(bold=True)
    _cel(ws, r + 2, 6, f"=COUNT(F{r0 + 1}:F{ult})")
    ws.cell(row=r + 4, column=1, value="Fatores: índices de correção dos salários de contribuição (INPC; IGP-DI de "
                                       "05/1996 a 08/2006 pela tabela do Manual CJF). Confira com a carta de concessão.")
    _impressao(ws, f"A{r0 + 1}")


def _sem_salario(ws, caso, ben):
    from motor import competencias_sem_salario
    r0 = _identificacao(ws, caso, "Competências Sem Salários de Contribuição", None)
    _cab(ws, r0, ["Competência", "Empresa", "Data"], [13, 50, 13])
    lista = caso.get("competencias_sem_salario") or competencias_sem_salario(caso)
    for i, x in enumerate(lista, r0 + 1):
        _cel(ws, i, 1, x["competencia"])
        _cel(ws, i, 2, x.get("empresa", ""))
        _cel(ws, i, 3, x.get("data", ""))


def _resumo_beneficio(ws, caso, ben):
    _identificacao(ws, caso, "Resumo do Cálculo do Benefício", None)
    if not ben:
        return
    melhor = next((x for x in ben["regras"] if x["regra"] == ben["melhor_regra"]), None)
    r = 8
    for rot, v in [("Tempo até a DER", ben["tempo_ate_der"]), ("Carência", ben["carencia_der"]),
                   ("Regra mais vantajosa", ben["melhor_regra"]),
                   ("DIB", melhor.get("dib") if melhor else ""), ("RMI", melhor.get("rmi") if melhor else "")]:
        ws.cell(row=r, column=1, value=rot).font = Font(bold=True)
        _cel(ws, r, 2, v, MOEDA if rot == "RMI" else None)
        r += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50


# --------------------------------------------------------------------------- liquidação

def _evolucao(ws, caso, liq, P):
    r0 = _identificacao(ws, caso, "Relatório da Evolução da R.M.I.", liq)
    _cab(ws, r0, ["Data", "Valor do Benefício", "Reajuste (%)", "Teto do Benefício", "Salário Mínimo",
                  "Benefício Mensal", "Fração (pró-rata)", "Benefício Considerado"],
         [11, 17, 13, 16, 14, 17, 12, 19])
    rows = {}
    r = r0 + 1
    for i, e in enumerate(liq["evolucao"]):
        _cel(ws, r, 1, e["competencia"])
        _cel(ws, r, 2, f"={P['rmi']}" if i == 0 else f"=F{r - 1}", MOEDA)
        _cel(ws, r, 3, e["reajuste"], "0.000000", fill=ENTRADA)
        _cel(ws, r, 4, e["teto"], MOEDA)
        _cel(ws, r, 5, e["sm"], MOEDA)
        teto = f"MIN(D{r}," if e["teto"] else "("
        _cel(ws, r, 6, f"=MAX(E{r},{teto}IF(C{r}=\"\",B{r},TRUNC(B{r}*(1+C{r}/100),2))))", MOEDA)
        _cel(ws, r, 7, e["fracao"], "0.0000")
        _cel(ws, r, 8, f"=TRUNC(F{r}*G{r},2)", MOEDA)
        rows[e["c"]] = r
        r += 1
    ws.cell(row=r + 1, column=1, value="Nova R.M.A.:").font = Font(bold=True)
    _cel(ws, r + 1, 2, f"=F{r - 1}", MOEDA, bold=True, fill=TOT_FILL)
    ws.cell(row=r + 2, column=1, value="Pró-rata no mês inicial (data da DIB ou da prescrição) e no mês final, "
                                       "proporcional ao número de dias (mês comercial de 30 dias). Valores truncados "
                                       "em centavos. 1º reajuste proporcional ao INPC do mês da DIB a dezembro.")
    _impressao(ws, f"A{r0 + 1}")
    return rows


def _diferencas(ws, caso, liq, P, ev):
    r0 = _identificacao(ws, caso, "Relatório das Diferenças Não Recebidas", liq)
    crit = liq["criterio"]
    notas = [
        f"Correção monetária: Tabela da Justiça Federal ({crit}; IGP-DI de 05/1996 a 08/2006), pró-rata no 1º mês, até 11/2021.",
        "Juros de mora a partir da citação: 1% a.m. até 06/2009; 0,5% a.m. até 04/2012; juros da poupança "
        "(Lei 12.703/2012) até 11/2021 — sem capitalização.",
        "De 12/2021 a 08/2025: SELIC acumulada (simples), uma única vez, substituindo correção e juros "
        "(EC 113/2021, art. 3º; Tema 1.419/STF).",
    ]
    if liq["fase_c"]:
        notas.append("A partir de 09/2025: correção pelo INPC e juros de mora pela taxa legal (SELIC com dedução do "
                     "INPC; art. 406 do CC) — Manual de Cálculos da JF, Res. CJF 990/2026 (EC 136/2025).")
    notas.append(f"Diferenças consideradas de {d_fmt(liq['inicio'])} a {d_fmt(liq['fim'])}.")
    for i, n in enumerate(notas):
        ws.cell(row=r0 + i, column=1, value=n).font = Font(italic=True, size=9)
    h = r0 + len(notas) + 1
    cols = ["Data", "Valor do Benefício", "Valor Já Recebido", "Diferença", "Índice de Atualização",
            "Valor Atualizado", "% Juros", "Valor dos Juros", "Valor Atualizado + Juros", "% Selic",
            "Valor da Selic", "Total até 08/2025" if liq["fase_c"] else "Total Atualizado"]
    larg = [9, 15, 13, 14, 12, 14, 9, 13, 15, 9, 13, 15]
    if liq["fase_c"]:
        cols += ["Índice INPC desde 09/2025", "Valor Corrigido", "% Taxa Legal", "Juros Taxa Legal", "Total Atualizado"]
        larg += [12, 14, 10, 13, 15]
    cols += ["Base atualizada", "Juros (todos)", "Ano", "Tipo"]
    larg += [14, 14, 7, 8]
    _cab(ws, h, cols, larg)
    n = len(cols)
    C = {nome: L(i) for i, nome in enumerate(
        ["data", "ben", "rec", "dif", "idx", "atu", "jp", "jv", "aj", "sp", "sv", "tb"] +
        (["ic", "vc", "tlp", "tlv", "tot"] if liq["fase_c"] else []) + ["base", "juros", "ano", "tipo"], 1)}
    if not liq["fase_c"]:
        C["tot"] = C["tb"]
    r = h + 1
    primeira = r
    rows = []
    for ln in liq["linhas"]:
        f = ln["formula"]
        if f[0] == "mes":
            ben = f"='Evolução RMI'!H{ev[f[1]]}"
        elif f[0] == "13a":
            ben = f"=TRUNC('Evolução RMI'!F{ev[f[1]]}/2,2)"
        elif f[0] == "13u":
            ben = f"=TRUNC('Evolução RMI'!F{ev[f[1]]}*{f[2]}/12,2)"
        else:
            ben = f"=TRUNC('Evolução RMI'!F{ev[f[1]]}*{f[2]}/12-'Evolução RMI'!F{ev[f[3]]}/2,2)"
        v = {
            "data": ln["rotulo"], "ben": ben, "rec": ln["recebido"], "dif": f"={C['ben']}{r}-{C['rec']}{r}",
            "idx": ln["indice"], "atu": f"=ROUND({C['dif']}{r}*{C['idx']}{r},2)", "jp": ln["juros_pct"],
            "jv": f"=ROUND({C['atu']}{r}*{C['jp']}{r}/100,2)", "aj": f"={C['atu']}{r}+{C['jv']}{r}",
            "sp": ln["selic_pct"], "sv": f"=ROUND({C['aj']}{r}*{C['sp']}{r}/100,2)",
            "tb": f"={C['aj']}{r}+{C['sv']}{r}", "ano": ln["c"][0], "tipo": "13º" if ln["tipo"] == "13" else "mensal",
        }
        if liq["fase_c"]:
            v.update({"ic": ln["indice_c"], "vc": f"=ROUND({C['tb']}{r}*{C['ic']}{r},2)", "tlp": ln["taxa_legal_pct"],
                      "tlv": f"=ROUND({C['vc']}{r}*{C['tlp']}{r}/100,2)", "tot": f"={C['vc']}{r}+{C['tlv']}{r}"})
            v["base"] = f"={C['atu']}{r}+{C['vc']}{r}-{C['tb']}{r}"
            v["juros"] = f"={C['jv']}{r}+{C['sv']}{r}+{C['tlv']}{r}"
        else:
            v["base"] = f"={C['atu']}{r}"
            v["juros"] = f"={C['jv']}{r}+{C['sv']}{r}"
        fmts = {"idx": "0.000000", "ic": "0.000000", "jp": "0.0000", "sp": "0.0000", "tlp": "0.0000",
                "ano": "0", "data": None, "tipo": None}
        for nome, col in C.items():
            if nome == "tot" and not liq["fase_c"]:
                continue
            editavel = nome in ("rec", "idx", "jp", "sp", "ic", "tlp")
            _cel(ws, r, ws[col + "1"].column, v[nome], fmts.get(nome, MOEDA), fill=ENTRADA if editavel else None)
        rows.append((ln, r))
        r += 1
    ultima = r - 1
    ws.cell(row=r, column=1, value="Totais:").font = Font(bold=True)
    for nome in ["ben", "rec", "dif", "atu", "jv", "aj", "sv", "tb"] + (["vc", "tlv", "tot"] if liq["fase_c"] else []) + ["base", "juros"]:
        col = C[nome]
        _cel(ws, r, ws[col + "1"].column, f"=SUM({col}{primeira}:{col}{ultima})", MOEDA, bold=True, fill=TOT_FILL)
    tot_row = r
    ws.cell(row=r + 2, column=1, value="R.M.A.:").font = Font(bold=True)
    _cel(ws, r + 2, 2, f"='Evolução RMI'!F{max(ev.values())}", MOEDA)
    ws.cell(row=r + 3, column=1, value=f"Nº de meses com diferenças: {sum(1 for x in liq['linhas'] if x['tipo'] == 'mensal')}")
    ws.cell(row=r + 4, column=1, value="Tabela da Justiça Federal (previdenciário): ORTN 10/1964–02/1986; OTN 03/1986–01/1989; "
                                       "IPC 01/1989–02/1989; BTN 03/1989–03/1990; IPC 03/1990–02/1991; INPC 03/1991–12/1992; "
                                       "IRSM 01/1993–02/1994; URV 03/1994–06/1994; IPC-r 07/1994–06/1995; INPC 07/1995–04/1996; "
                                       "IGP-DI 05/1996–08/2006; INPC a partir de 09/2006.").font = Font(size=9)
    adv = caso.get("advogado", {})
    if adv.get("nome"):
        ws.cell(row=r + 6, column=1, value=f"{adv['nome']}  {adv.get('oab', '')}").font = Font(bold=True)
    _impressao(ws, f"B{h + 1}")
    return dict(C=C, primeira=primeira, ultima=ultima, tot=tot_row, rows=rows)


def _honorarios(ws, caso, liq, P, dif):
    r0 = _identificacao(ws, caso, "Cálculo dos Honorários", liq)
    ws.cell(row=r0, column=1, value="Honorários de").font = Font(bold=True)
    _cel(ws, r0, 2, f"={P['pct']}", "0.0000")
    ws.cell(row=r0, column=3, value="% sobre a base de cálculo atualizada e sobre os juros.")
    h = r0 + 2
    _cab(ws, h, ["Data", "Base de Cálculo", "Base de Cálculo Atualizada", "Juros Sobre a Base de Cálculo",
                 "Base Atualizada + Juros", "Valor Calculado"], [9, 16, 18, 18, 18, 16])
    C = dif["C"]
    hon_ids = {id(x) for x in liq["hon_linhas"]}
    r = h + 1
    primeira = r
    for ln, dr in dif["rows"]:
        if id(ln) not in hon_ids:
            continue
        _cel(ws, r, 1, ln["rotulo"])
        _cel(ws, r, 2, f"='Diferenças'!{C['dif']}{dr}", MOEDA)
        _cel(ws, r, 3, f"='Diferenças'!{C['base']}{dr}", MOEDA)
        _cel(ws, r, 4, f"='Diferenças'!{C['juros']}{dr}", MOEDA)
        _cel(ws, r, 5, f"=C{r}+D{r}", MOEDA)
        _cel(ws, r, 6, f"=ROUND(E{r}*{P['pct']}/100,2)", MOEDA)
        r += 1
    ultima = r - 1
    ws.cell(row=r, column=1, value="Totais:").font = Font(bold=True)
    for col in "BCDEF":
        _cel(ws, r, ws[col + "1"].column, f"=SUM({col}{primeira}:{col}{ultima})", MOEDA, bold=True, fill=TOT_FILL)
    t = r
    linhas = [("Sobre a Base Atualizada:", f"=ROUND(C{t}*{P['pct']}/100,2)"),
              ("Sobre os Juros:", f"=ROUND(D{t}*{P['pct']}/100,2)"),
              ("Valor Fixo:", f"={P['hon_fixo']}"),
              ("Total Geral de Honorários:", f"=B{t + 2}+B{t + 3}+B{t + 4}")]
    for i, (rot, fml) in enumerate(linhas, t + 2):
        ws.cell(row=i, column=1, value=rot).font = Font(bold=True)
        _cel(ws, i, 2, fml, MOEDA, bold=i == t + 5, fill=TOT_FILL if i == t + 5 else None)
    ws.column_dimensions["A"].width = 26
    _impressao(ws, f"A{h + 1}")
    return dict(base=f"'Honorários'!$B${t + 2}", juros=f"'Honorários'!$B${t + 3}", fixo=f"'Honorários'!$B${t + 4}",
                total=f"'Honorários'!$B${t + 5}")


def _ir(ws, caso, liq, dif):
    r0 = _identificacao(ws, caso, "Informações para Imposto de Renda (Rendimentos Recebidos Acumuladamente)", liq)
    C = dif["C"]
    rng_tot = f"'Diferenças'!${C['tot']}${dif['primeira']}:${C['tot']}${dif['ultima']}"
    rng_ano = f"'Diferenças'!${C['ano']}${dif['primeira']}:${C['ano']}${dif['ultima']}"
    ano = liq["atual"].year
    _cab(ws, r0, ["", "Valor", "Meses"], [34, 18, 10])
    dados = [("Anos-base anteriores", f"=SUMIFS({rng_tot},{rng_ano},\"<{ano}\")", liq["ir"]["anteriores"][1]),
             (f"Ano-base atual ({ano})", f"=SUMIFS({rng_tot},{rng_ano},{ano})", liq["ir"]["atual"][1]),
             ("Total", f"=B{r0 + 1}+B{r0 + 2}", liq["ir"]["anteriores"][1] + liq["ir"]["atual"][1])]
    for i, (rot, f, m) in enumerate(dados, r0 + 1):
        _cel(ws, i, 1, rot, bold=True)
        _cel(ws, i, 2, f, MOEDA)
        _cel(ws, i, 3, m)
    ws.cell(row=r0 + 5, column=1, value="Meses: competências mensais + 1 mês por ano com 13º (art. 12-A da Lei 7.713/88).")


def _resumo(ws, caso, liq, P, dif, hon):
    r0 = _identificacao(ws, caso, "Resumo do Cálculo", liq)
    C, t = dif["C"], dif["tot"]
    D = lambda col: f"'Diferenças'!{C[col]}{t}"
    r = r0 + 1
    ws.cell(row=r0, column=1, value=f"Renda Mensal Corrigida até {d_fmt(liq['atual'])}:").font = Font(bold=True)
    _cel(ws, r0, 2, f"='Diferenças'!B{t + 2}", MOEDA, bold=True)
    itens = [
        ("principal", "Principal:", f"={D('dif')}"),
        ("correcao", "Correção Monetária:", f"={D('base')}-{D('dif')}"),
        ("atualizado", "Total Atualizado:", f"={D('base')}"),
        ("juros", "Valor dos Juros:", f"={D('juros')}"),
        ("acordo", "(-) Acordo/Deságio:", f"={P['acordo']}"),
        ("tj", "Total Atualizado + Valor dos Juros - Valor Acordo:", None),
        ("hb", "Honorários Sobre a Base de Cálculo Atualizada:", f"={hon['base']}"),
        ("hj", "Honorários Sobre os Juros da Base de Cálculo:", f"={hon['juros']}"),
        ("hf", "Valor Fixo de Honorários:", f"={hon['fixo']}"),
        ("ht", "Total Atualizado de Honorários:", f"={hon['total']}"),
        ("multa", "Multa:", f"={P['multa']}"),
        ("dano", "Dano Moral/Material:", f"={P['dano']}"),
        ("tg", "Total Geral:", None),
    ]
    pos = {}
    for i, (k, rot, f) in enumerate(itens):
        pos[k] = r + i
    for k, rot, f in itens:
        i = pos[k]
        if k == "tj":
            f = f"=B{pos['atualizado']}+B{pos['juros']}-B{pos['acordo']}"
        if k == "tg":
            f = f"=B{pos['tj']}+B{pos['ht']}+B{pos['multa']}+B{pos['dano']}"
        ws.cell(row=i, column=1, value=rot).font = Font(bold=True)
        _cel(ws, i, 2, f, MOEDA, bold=k in ("tj", "tg"), fill=TOT_FILL if k in ("tj", "tg") else None)
    r = pos["tg"] + 1
    ws.cell(row=r, column=1, value=f"Total geral atualizado até {d_fmt(liq['atual'])}")
    ws.cell(row=r + 1, column=1, value="O valor equivale a (salários mínimos):").font = Font(bold=True)
    _cel(ws, r + 1, 2, f"=ROUND(B{pos['tg']}/{P['sm']},2)", "0.00")
    r += 3
    ws.cell(row=r, column=1, value="Valor da causa (art. 292, §§ 1º e 2º, do CPC)").font = Font(bold=True, color=AZUL)
    vc = [("Parcelas vencidas (atualizadas + juros):", f"=B{pos['tj']}", MOEDA),
          ("12 parcelas vincendas (12 × R.M.A.):", f"=ROUND(12*'Diferenças'!B{t + 2},2)", MOEDA),
          ("Valor da causa:", f"=B{r + 1}+B{r + 2}", MOEDA),
          ("Limite do JEF (60 salários mínimos):", f"=ROUND(60*{P['sm']},2)", MOEDA),
          ("Competência:", f"=IF(B{r + 3}<=B{r + 4},\"Juizado Especial Federal\",\"Vara Federal (ou renúncia ao excedente)\")", None)]
    for i, (rot, f, fmt) in enumerate(vc, r + 1):
        ws.cell(row=i, column=1, value=rot).font = Font(bold=True)
        _cel(ws, i, 2, f, fmt)
    adv = caso.get("advogado", {})
    if adv.get("nome"):
        ws.cell(row=r + 8, column=1, value=f"{adv['nome']}  {adv.get('oab', '')}").font = Font(bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 28
    _impressao(ws)


def _indices(ws, series):
    ws["A1"] = "Índices mensais utilizados (%)"
    ws["A1"].font = TITULO
    nomes = [n for n in ("inpc", "igpdi", "ipcae", "juros_mora", "selic", "taxa_legal") if series.usadas.get(n)]
    _cab(ws, 3, ["Competência"] + [n.upper().replace("_", " ") for n in nomes], [13] + [13] * len(nomes))
    meses = sorted({k for n in nomes for k in series.usadas[n]})
    for i, k in enumerate(meses, 4):
        a, m = k.split("-")
        _cel(ws, i, 1, f"{m}/{a}")
        for j, n in enumerate(nomes, 2):
            _cel(ws, i, j, series.usadas[n].get(k), "0.0000")
    ws.cell(row=len(meses) + 5, column=1, value="Fontes: IBGE/BCB-SGS (INPC 188, IGP-DI 190, IPCA-E 10764, SELIC 4390, "
                                                  "meta SELIC 432) e Manual de Cálculos da Justiça Federal.")
